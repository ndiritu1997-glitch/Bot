"""Readers for Word documents and plain-text / code source files.

These functions only READ source content. They never modify or save over the
source file.
"""

from __future__ import annotations

from pathlib import Path
from typing import List


class SourceError(Exception):
    """Raised when a source file is missing, empty, or unreadable."""


def read_code_file(path: Path) -> str:
    """Read an HTML/CSS/JS/PHP (or any text) source file exactly as-is.

    Preserves indentation, blank lines, punctuation, and character order.
    """
    if not path.exists():
        raise SourceError(f"Input file not found: {path}")
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Fall back to latin-1 so odd bytes don't crash the read.
        text = path.read_text(encoding="latin-1")
    except OSError as exc:
        raise SourceError(f"Could not read {path}: {exc}") from exc

    if text.strip() == "":
        raise SourceError(f"Input file is empty: {path}")
    return text


def read_docx_text(path: Path) -> str:
    """Extract paragraph text from a .docx file, preserving paragraph order.

    Returns plain text with paragraphs separated by newlines. This first
    version handles body paragraphs only. It does NOT reproduce tables,
    images, headers/footers, footnotes, styles, or embedded objects — callers
    should warn the user about that (see detect_unsupported_docx).
    """
    if not path.exists():
        raise SourceError(f"Input file not found: {path}")

    try:
        import docx  # python-docx  # type: ignore
    except ImportError as exc:
        raise SourceError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises various errors on bad files
        raise SourceError(f"Could not open Word document {path}: {exc}") from exc

    paragraphs: List[str] = [p.text for p in document.paragraphs]
    text = "\n".join(paragraphs)

    if text.strip() == "":
        raise SourceError(
            f"No typable paragraph text found in {path}. "
            f"(Tables/images are not supported in this version.)"
        )
    return text


def detect_unsupported_docx(path: Path) -> List[str]:
    """Return a list of warnings about content that will NOT be reproduced."""
    warnings: List[str] = []
    try:
        import docx  # type: ignore

        document = docx.Document(str(path))
    except Exception:
        return warnings  # reading issues are handled elsewhere

    if document.tables:
        warnings.append(
            f"{len(document.tables)} table(s) found - table content will NOT be typed."
        )
    # Inline images show up as drawing elements in the XML.
    try:
        body_xml = document.element.body.xml
        if "w:drawing" in body_xml or "graphicData" in body_xml:
            warnings.append("Image(s)/drawing(s) found - these will NOT be reproduced.")
    except Exception:
        pass
    return warnings
