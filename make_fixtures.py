"""Generate the two binary input fixtures: inputdocx.docx and inputppt.pptx.

The text fixtures (html/css/js/php/csv) ship ready-made in inputs/. The Word
and PowerPoint samples are binary, so this script builds them locally. Run it
once after installing requirements:

    python make_fixtures.py

It writes into the inputs/ folder next to this script and never overwrites a
file that already exists (delete the file first if you want it regenerated).
You can also just drop your own inputdocx.docx / inputppt.pptx into inputs/.
"""

from __future__ import annotations

from pathlib import Path

INPUT_DIR = Path(__file__).resolve().parent / "inputs"


def make_docx(path: Path) -> None:
    if path.exists():
        print(f"skip (exists): {path}")
        return
    import docx

    doc = docx.Document()
    doc.add_paragraph("Typing Simulation Sample Document")
    doc.add_paragraph(
        "This is the first body paragraph. The bot reads these paragraphs and "
        "types them one character at a time into Microsoft Word."
    )
    doc.add_paragraph(
        "This is a second paragraph. Punctuation, such as commas, periods, and "
        "question marks, is preserved. Does it pause after sentences? Yes."
    )
    doc.add_paragraph("A short closing line.")
    doc.save(str(path))
    print(f"wrote: {path}")


def make_pptx(path: Path) -> None:
    if path.exists():
        print(f"skip (exists): {path}")
        return
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Typing Simulation Sample Deck"
    s1.placeholders[1].text = "Generated as a test fixture"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "What this deck is for"
    s2.placeholders[1].text = (
        "Testing keystroke recording\n"
        "Slide text is typed, not pasted\n"
        "Each slide is recreated in order"
    )

    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Second content slide"
    s3.placeholders[1].text = "Point one\nPoint two\nPoint three"

    prs.save(str(path))
    print(f"wrote: {path}")


def main() -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    make_docx(INPUT_DIR / "inputdocx.docx")
    make_pptx(INPUT_DIR / "inputppt.pptx")
    print("Done.")


if __name__ == "__main__":
    main()
